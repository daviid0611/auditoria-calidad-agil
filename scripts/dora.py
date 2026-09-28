"""Métricas DORA a partir de datos/despliegues.csv (bloque 5 del taller).

Uso, desde la raíz del repositorio:
    python scripts/dora.py          imprime las métricas y el análisis por día
    python scripts/dora.py --xlsx   además genera datos/metricas_dora.xlsx

El script está fuera de src/ para que no cuente en la cobertura del CI.
El cálculo solo usa la biblioteca estándar; openpyxl se necesita únicamente
para --xlsx (pip install openpyxl).
"""

import argparse
import csv
import statistics
import sys
from datetime import datetime
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
RUTA_CSV = RAIZ / "datos" / "despliegues.csv"
RUTA_XLSX = RAIZ / "datos" / "metricas_dora.xlsx"
PERIODO_DIAS = 28  # Dato del enunciado (datos/LEEME.md).
FORMATO_FECHA = "%Y-%m-%d %H:%M"
DIAS_SEMANA = (
    "lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"
)
VALORES_CONTROL = {
    "lead_time_mediana_h": 20.0,
    "tasa_fallo": 0.20,
    "recuperacion_promedio_h": 4.5,
}


def leer_despliegues(ruta: Path = RUTA_CSV) -> list[dict]:
    """Lee el CSV y convierte fechas, resultado y horas a tipos de Python.

    Lanza ValueError si la columna exitoso tiene un valor distinto de si/no.
    """
    with open(ruta, newline="", encoding="utf-8") as archivo:
        filas = list(csv.DictReader(archivo))
    for fila in filas:
        if fila["exitoso"] not in ("si", "no"):
            raise ValueError(f"exitoso no válido en id {fila['id']}")
        fila["fecha_commit"] = datetime.strptime(
            fila["fecha_commit"], FORMATO_FECHA
        )
        fila["fecha_despliegue"] = datetime.strptime(
            fila["fecha_despliegue"], FORMATO_FECHA
        )
        fila["fallido"] = fila["exitoso"] == "no"
        fila["horas_recuperacion"] = float(fila["horas_recuperacion"])
    return filas


def lead_time_horas(despliegue: dict) -> float:
    """Horas entre el commit y el despliegue."""
    diferencia = despliegue["fecha_despliegue"] - despliegue["fecha_commit"]
    return diferencia.total_seconds() / 3600


def calcular_metricas(
    despliegues: list[dict], periodo_dias: int = PERIODO_DIAS
) -> dict:
    """Calcula las 4 métricas DORA.

    - Frecuencia: despliegues / periodo_dias (y por semana).
    - Lead time: mediana de (despliegue - commit) en horas.
    - Tasa de fallo: fallidos / total.
    - Recuperación: promedio de horas_recuperacion de los fallidos.
    """
    total = len(despliegues)
    fallidos = [d for d in despliegues if d["fallido"]]
    horas_fallidos = [d["horas_recuperacion"] for d in fallidos]
    return {
        "total": total,
        "fallidos": len(fallidos),
        "frecuencia_por_dia": total / periodo_dias,
        "frecuencia_por_semana": total / periodo_dias * 7,
        "lead_time_mediana_h": statistics.median(
            lead_time_horas(d) for d in despliegues
        ),
        "tasa_fallo": len(fallidos) / total,
        "recuperacion_promedio_h": (
            statistics.mean(horas_fallidos) if horas_fallidos else 0.0
        ),
    }


def dia_semana(fecha: datetime) -> str:
    """Nombre del día de la semana en español."""
    return DIAS_SEMANA[fecha.weekday()]


def fallos_por_dia(despliegues: list[dict]) -> list[tuple[str, int, int]]:
    """(día, despliegues, fallidos) por día de la semana del despliegue."""
    resumen = []
    for dia in DIAS_SEMANA:
        del_dia = [
            d for d in despliegues if dia_semana(d["fecha_despliegue"]) == dia
        ]
        resumen.append((dia, len(del_dia), sum(d["fallido"] for d in del_dia)))
    return resumen


def numero(valor: float, decimales: int = 2) -> str:
    """Número con coma decimal, como se escribe en español."""
    return f"{valor:.{decimales}f}".replace(".", ",")


def imprimir_reporte(despliegues: list[dict]) -> None:
    """Imprime las métricas, la comparación con los valores de control
    y el análisis de fallos por día de la semana."""
    m = calcular_metricas(despliegues)
    print(f"Métricas DORA ({m['total']} despliegues, {PERIODO_DIAS} días)")
    print(
        f"  Frecuencia de despliegue: {numero(m['frecuencia_por_dia'])} "
        f"por día ({numero(m['frecuencia_por_semana'], 1)} por semana)"
    )
    print(f"  Lead time, mediana: {numero(m['lead_time_mediana_h'], 1)} h")
    print(
        f"  Tasa de fallo: {m['fallidos']}/{m['total']} = "
        f"{numero(m['tasa_fallo'] * 100, 1)} %"
    )
    print(
        f"  Recuperación, promedio de fallidos: "
        f"{numero(m['recuperacion_promedio_h'], 1)} h"
    )

    print("\nValores de control")
    for clave, esperado in VALORES_CONTROL.items():
        ok = "sí" if abs(m[clave] - esperado) < 1e-9 else "NO"
        print(f"  {clave}: esperado {esperado}, obtenido {m[clave]} -> {ok}")

    print("\nFallos por día de la semana del despliegue")
    for dia, total, fallidos in fallos_por_dia(despliegues):
        tasa = numero(fallidos / total * 100, 0) + " %" if total else "-"
        print(f"  {dia:<10} {total:>2} despliegues, {fallidos} fallidos ({tasa})")

    print("\nDespliegues fallidos")
    for d in despliegues:
        if d["fallido"]:
            print(
                f"  id {d['id']:>2}: commit {dia_semana(d['fecha_commit'])} "
                f"{d['fecha_commit']:%Y-%m-%d %H:%M}, despliegue "
                f"{dia_semana(d['fecha_despliegue'])} "
                f"{d['fecha_despliegue']:%Y-%m-%d %H:%M}, "
                f"recuperación {numero(d['horas_recuperacion'], 0)} h"
            )


def generar_xlsx(despliegues: list[dict], ruta: Path = RUTA_XLSX) -> None:
    """Genera la hoja de cálculo con fórmulas (no valores pegados).

    Hojas: despliegues (datos + columnas calculadas), metricas y por_dia.
    Las celdas en azul son datos de entrada; las negras son fórmulas.
    """
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill

    fuente = Font(name="Arial", size=10)
    entrada = Font(name="Arial", size=10, color="0000FF")
    titulo = Font(name="Arial", size=10, bold=True)
    relleno = PatternFill("solid", start_color="DDEBF7")
    dias = ",".join(f'"{dia}"' for dia in DIAS_SEMANA)
    n = len(despliegues) + 1  # última fila con datos

    libro = Workbook()

    # Hoja 1: datos originales + columnas calculadas con fórmula.
    hoja = libro.active
    hoja.title = "despliegues"
    encabezados = [
        "id", "fecha_commit", "fecha_despliegue", "exitoso",
        "horas_recuperacion", "lead_time_h", "dia_despliegue",
        "dia_commit", "fallido",
    ]
    hoja.append(encabezados)
    for i, d in enumerate(despliegues, start=2):
        hoja.append([
            int(d["id"]), d["fecha_commit"], d["fecha_despliegue"],
            d["exitoso"], d["horas_recuperacion"],
            f"=ROUND((C{i}-B{i})*24,2)",
            f"=CHOOSE(WEEKDAY(C{i},2),{dias})",
            f"=CHOOSE(WEEKDAY(B{i},2),{dias})",
            f'=IF(D{i}="no",1,0)',
        ])
    for fila in hoja.iter_rows(min_row=2, max_row=n):
        for celda in fila:
            celda.font = entrada if celda.column <= 5 else fuente
        fila[1].number_format = "yyyy-mm-dd hh:mm"
        fila[2].number_format = "yyyy-mm-dd hh:mm"
    for ancho, columna in zip(
        (5, 17, 17, 9, 19, 12, 16, 12, 9), "ABCDEFGHI"
    ):
        hoja.column_dimensions[columna].width = ancho

    # Hoja 2: las 4 métricas DORA.
    met = libro.create_sheet("metricas")
    rango = f"despliegues!$D$2:$D${n}"

    def coincide(fila: int) -> str:
        return f'=IF(ROUND(B{fila},6)=ROUND(D{fila},6),"sí","no")'

    filas_metricas = [
        ("Métricas DORA (datos/despliegues.csv)", None, None, None, None),
        (None, None, None, None, None),
        ("Periodo (días)", PERIODO_DIAS, "Dato del enunciado", None, None),
        ("Total de despliegues", f"=COUNT(despliegues!$A$2:$A${n})",
         "Filas con id", None, None),
        ("Despliegues fallidos", f'=COUNTIF({rango},"no")',
         "exitoso = no", None, None),
        (None, None, None, None, None),
        ("Métrica", "Valor", "Fórmula", "Valor de control", "¿Coincide?"),
        ("Frecuencia de despliegue (por día)", "=B4/B3",
         "total / periodo", None, None),
        ("Frecuencia de despliegue (por semana)", "=B8*7",
         "por día × 7", None, None),
        ("Lead time de cambios, mediana (h)",
         f"=MEDIAN(despliegues!$F$2:$F${n})",
         "mediana de (despliegue - commit) en horas",
         VALORES_CONTROL["lead_time_mediana_h"], coincide(10)),
        ("Tasa de fallo de cambios", "=B5/B4", "fallidos / total",
         VALORES_CONTROL["tasa_fallo"], coincide(11)),
        ("Tiempo de recuperación, promedio (h)",
         f'=AVERAGEIF({rango},"no",despliegues!$E$2:$E${n})',
         "promedio de horas_recuperacion de los fallidos",
         VALORES_CONTROL["recuperacion_promedio_h"], coincide(12)),
        (None, None, None, None, None),
        ("Leyenda: azul = dato de entrada; negro = fórmula.",
         None, None, None, None),
    ]
    for fila in filas_metricas:
        met.append(fila)
    for fila in met.iter_rows():
        for celda in fila:
            celda.font = fuente
    met["A1"].font = titulo
    for celda in met[7]:
        celda.font = titulo
        celda.fill = relleno
    met["B3"].font = entrada
    met["B3"].comment = Comment("Periodo indicado en datos/LEEME.md.", "Equipo")
    for fila in (10, 11, 12):
        met[f"D{fila}"].font = entrada
    met["D10"].comment = Comment(
        "Valores de control del enunciado del taller.", "Equipo"
    )
    met["B8"].number_format = "0.00"
    met["B9"].number_format = "0.0"
    met["B10"].number_format = "0.0"
    met["B11"].number_format = "0.0%"
    met["D11"].number_format = "0.0%"
    met["B12"].number_format = "0.0"
    for ancho, columna in zip((40, 10, 46, 16, 11), "ABCDE"):
        met.column_dimensions[columna].width = ancho

    # Hoja 3: fallos por día de la semana del despliegue.
    pdia = libro.create_sheet("por_dia")
    pdia.append(["Día del despliegue", "Despliegues", "Fallidos", "Tasa de fallo"])
    dia_col = f"despliegues!$G$2:$G${n}"
    for i, dia in enumerate(DIAS_SEMANA, start=2):
        pdia.append([
            dia,
            f"=COUNTIF({dia_col},A{i})",
            f'=COUNTIFS({dia_col},A{i},{rango},"no")',
            f"=IF(B{i}=0,0,C{i}/B{i})",
        ])
    pdia.append(["Total", "=SUM(B2:B8)", "=SUM(C2:C8)", "=IF(B9=0,0,C9/B9)"])
    pdia.append([None])
    pdia.append([
        "Lunes a jueves (ventana de la política)",
        "=SUM(B2:B5)", "=SUM(C2:C5)", "=IF(B11=0,0,C11/B11)",
    ])
    pdia.append([
        "Viernes a domingo (fuera de la ventana)",
        "=SUM(B6:B8)", "=SUM(C6:C8)", "=IF(B12=0,0,C12/B12)",
    ])
    for fila in pdia.iter_rows():
        for celda in fila:
            celda.font = fuente
        fila[3].number_format = "0%"
    for celda in pdia[1]:
        celda.font = titulo
        celda.fill = relleno
    pdia["A9"].font = titulo
    pdia.column_dimensions["A"].width = 40
    for columna in "BCD":
        pdia.column_dimensions[columna].width = 13

    libro.save(ruta)
    print(f"\nHoja de cálculo generada: {ruta.relative_to(RAIZ).as_posix()}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "--xlsx", action="store_true",
        help="genera datos/metricas_dora.xlsx con fórmulas",
    )
    args = parser.parse_args()
    # Tildes correctas también al redirigir la salida en Windows.
    sys.stdout.reconfigure(encoding="utf-8")
    despliegues = leer_despliegues()
    imprimir_reporte(despliegues)
    if args.xlsx:
        generar_xlsx(despliegues)


if __name__ == "__main__":
    main()
